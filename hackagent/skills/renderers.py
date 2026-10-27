"""Turn model-written content into files, locally and without model calls.

The model writes text (markdown, CSV, a JSON slide outline, source code);
these renderers convert it into the requested format. Optional libraries
are imported lazily so the intake jobs do not need them installed.
"""

from __future__ import annotations

import csv
import io
import json
import re

CONTENT_TYPES = {
    ".md": "text/markdown",
    ".txt": "text/plain",
    ".mmd": "text/plain",
    ".csv": "text/csv",
    ".json": "application/json",
    ".py": "text/x-python",
    ".js": "text/javascript",
    ".ts": "text/plain",
    ".tsx": "text/plain",
    ".html": "text/plain",  # served as plain text: never rendered from storage
    ".css": "text/css",
    ".toml": "text/plain",
    ".yaml": "text/plain",
    ".yml": "text/plain",
    ".sql": "text/plain",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pdf": "application/pdf",
}


def _md_blocks(md: str) -> list[tuple[str, str]]:
    """Very small markdown reader: (kind, text) with kind h1-h3, bullet, para."""
    out = []
    for line in md.splitlines():
        s = line.strip()
        if not s:
            continue
        if m := re.match(r"^(#{1,3})\s+(.*)", s):
            out.append((f"h{len(m.group(1))}", m.group(2)))
        elif m := re.match(r"^[-*]\s+(.*)", s):
            out.append(("bullet", m.group(1)))
        else:
            out.append(("para", s))
    return out


def _strip_inline(text: str) -> str:
    return re.sub(r"[*_`]{1,3}([^*_`]+)[*_`]{1,3}", r"\1", text)


def to_docx(md: str) -> bytes:
    from docx import Document

    doc = Document()
    for kind, text in _md_blocks(md):
        text = _strip_inline(text)
        if kind.startswith("h"):
            doc.add_heading(text, level=int(kind[1]))
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(md: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    sizes = {"h1": 18, "h2": 15, "h3": 13}
    for kind, text in _md_blocks(md):
        # Core PDF fonts are Latin-1 only.
        text = _strip_inline(text).encode("latin-1", "replace").decode("latin-1")
        pdf.set_font("Helvetica", "B" if kind in sizes else "", sizes.get(kind, 11))
        if kind == "bullet":
            text = "- " + text
        pdf.multi_cell(0, 7 if kind not in sizes else 9, text)
        pdf.ln(1)
    return bytes(pdf.output())


def to_pptx(outline: str) -> bytes:
    from pptx import Presentation

    data = json.loads(outline)
    slides = data["slides"] if isinstance(data, dict) else data
    prs = Presentation()
    for s in slides[:40]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(s.get("title", ""))[:200]
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate((s.get("bullets") or [])[:10]):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(bullet)[:300]
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])[:2000]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")


def _neutralise(cell: str) -> str:
    """Prevent spreadsheet formula injection from model-written CSV."""
    if cell.startswith(_FORMULA_PREFIX):
        try:
            float(cell)
            return cell
        except ValueError:
            return "'" + cell
    return cell


def to_xlsx(csv_text: str) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in csv.reader(io.StringIO(csv_text)):
        ws.append([float(c) if re.fullmatch(r"-?\d+(\.\d+)?", c) else _neutralise(c) for c in row])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(csv_text: str) -> bytes:
    rows = [[_neutralise(c) for c in row] for row in csv.reader(io.StringIO(csv_text))]
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode()


def render(path: str, content: str) -> tuple[bytes, str]:
    """Render ``content`` for ``path`` by extension; returns (bytes, content type)."""
    ext = "." + path.rsplit(".", 1)[-1].lower()
    if ext not in CONTENT_TYPES:
        raise ValueError(f"No renderer for {ext}")
    if ext == ".docx":
        data = to_docx(content)
    elif ext == ".pdf":
        data = to_pdf(content)
    elif ext == ".pptx":
        data = to_pptx(content)
    elif ext == ".xlsx":
        data = to_xlsx(content)
    elif ext == ".csv":
        data = to_csv(content)
    else:
        data = content.encode()
    return data, CONTENT_TYPES[ext]
