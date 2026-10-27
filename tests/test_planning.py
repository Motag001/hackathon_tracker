from __future__ import annotations

import io
import json
from datetime import timedelta
from decimal import Decimal

import pytest

from hackagent.agents import notifier, planner
from hackagent.agents.planner import PlanDraft, license_ok
from hackagent.shared.context import pause_for_budget, resume_budget_paused
from hackagent.shared.cost import BudgetExceeded
from hackagent.skills import registry, renderers
from tests.helpers import FIXED_NOW, make_ctx


def plan_payload(**over):
    base = {
        "problem": "Farmers in Kano lack timely price data.",
        "angle": "SMS-first price alerts.",
        "summary": "Build a price alert prototype.",
        "local_context": "NBS food price data (source: nigerianstat.gov.ng)",
        "milestones": [
            {
                "title": "Scope",
                "description": "",
                "est_hours": 2,
                "skills": ["doc_markdown"],
                "input_keys": [],
                "needs_top_tier": False,
                "top_tier_reason": "",
            },
            {
                "title": "Build",
                "description": "",
                "est_hours": 10,
                "skills": ["code", "telepathy"],
                "input_keys": ["Farmer Story"],
                "needs_top_tier": True,
                "top_tier_reason": "Tricky forecasting",
            },
        ],
        "datasets": [
            {
                "name": "NBS prices",
                "source": "NBS",
                "url": None,
                "license": "CC BY 4.0",
                "why": "",
                "owner_data": False,
            },
            {
                "name": "Kaggle crops",
                "source": "Kaggle",
                "url": None,
                "license": "CC BY-NC 4.0",
                "why": "",
                "owner_data": False,
            },
        ],
        "inputs_needed": [],
        "disclosure_draft": "I used Claude (AI) to draft code; I made all final decisions.",
        "risks": [],
    }
    base.update(over)
    return base


def _event(ctx, status="approved"):
    return ctx.store.insert(
        "events",
        {
            "source_id": "devpost",
            "external_id": "1",
            "dedupe_key": "k",
            "title": "Agri Hack",
            "url": "https://agri.devpost.com/",
            "status": status,
            "submission_deadline": (FIXED_NOW + timedelta(days=20)).isoformat(),
        },
    )[0]


def test_license_checks():
    assert license_ok("CC BY 4.0") is True
    assert license_ok("CC0") is True
    assert license_ok("CC BY-NC-SA 4.0") is False
    assert license_ok("Custom research license") is None
    assert license_ok(None) is None


def test_validate_fixes_skills_inputs_and_flags_licenses():
    ctx = make_ctx()
    draft, warnings = planner.validate(
        PlanDraft(**plan_payload()), ctx, set(registry.load_skills())
    )
    assert draft.milestones[1].skills == ["code"]
    keys = {i.key for i in draft.inputs_needed}
    assert "farmer_story" in keys  # referenced by a milestone, so requested
    assert any(k.startswith("dataset_license_kaggle") for k in keys)
    assert not any(k.startswith("dataset_license_nbs") for k in keys)
    assert any("telepathy" in w for w in warnings)


def test_validate_warns_on_hours_and_disclosure():
    ctx = make_ctx()
    payload = plan_payload(disclosure_draft="Made by me.")
    payload["milestones"][1]["est_hours"] = 90
    _, warnings = planner.validate(PlanDraft(**payload), ctx, set(registry.load_skills()))
    assert any("over the 72h" in w for w in warnings)
    assert any("does not mention AI" in w for w in warnings)


def test_draft_plan_and_activate():
    ctx = make_ctx(lambda call: plan_payload())
    ev = _event(ctx)
    stats = planner.run(ctx)
    assert stats["planned"] == 1
    (plan,) = ctx.store.select("plans")
    assert plan["status"] == "ready" and plan["total_hours"] == 12
    assert ctx.store.select("events")[0]["status"] == "plan_ready"
    assert len(ctx.store.select("milestones")) == 2
    assert {d["license_ok"] for d in ctx.store.select("datasets")} == {True, False}
    call = ctx.llm.calls[0]
    assert call.tier.value == "mid" and call.priority.value == "essential"
    assert "Skill registry" in call.stable_context

    # The owner approves (the DB trigger does this in production).
    ctx.store.update("plans", {"id": plan["id"]}, {"status": "approved"})
    ctx.store.update("events", {"id": ev["id"]}, {"status": "plan_approved"})
    planner.run(ctx)
    kinds = {r["kind"] for r in ctx.store.select("input_requests")}
    assert kinds == {"other", "dataset", "top_tier_approval"}
    build = next(m for m in ctx.store.select("milestones") if m["seq"] == 2)
    assert "top_tier_2" in build["input_keys"]
    assert ctx.store.select("events")[0]["status"] == "building"


def test_revision_supersedes_previous_plan():
    ctx = make_ctx(lambda call: plan_payload())
    ev = _event(ctx, status="planning")
    ctx.store.insert(
        "plans",
        {
            "event_id": ev["id"],
            "version": 1,
            "status": "revision_requested",
            "summary": "old",
            "revision_notes": "Make it smaller",
        },
    )
    planner.run(ctx)
    plans = sorted(ctx.store.select("plans"), key=lambda p: p["version"])
    assert [p["status"] for p in plans] == ["superseded", "ready"]
    assert "Make it smaller" in ctx.llm.calls[0].prompt


def test_budget_pause_and_resume():
    ctx = make_ctx()
    ev = _event(ctx, status="building")
    pause_for_budget(ctx, ev, BudgetExceeded("event_budget", "Event budget $1.50 reached"))
    assert ctx.store.select("events")[0]["status"] == "paused_budget"
    assert resume_budget_paused(ctx) == 0  # still pending
    ctx.store.update(
        "input_requests", {"event_id": ev["id"]}, {"status": "answered", "response": "yes"}
    )
    assert resume_budget_paused(ctx) == 1
    assert ctx.store.select("events")[0]["status"] == "building"
    with pytest.raises(BudgetExceeded):
        pause_for_budget(ctx, ev, BudgetExceeded("monthly_cap", "cap"))


# -- renderers ---------------------------------------------------------------


def test_renderers_produce_valid_files():
    md = "# Title\n\nSome **bold** text.\n\n- one\n- two ₦\n"
    docx, ct = renderers.render("report.docx", md)
    assert docx[:2] == b"PK" and "wordprocessingml" in ct
    pdf, _ = renderers.render("report.pdf", md)
    assert pdf.startswith(b"%PDF")
    outline = json.dumps({"slides": [{"title": "Hi", "bullets": ["a", "b"], "notes": "n"}]})
    pptx, _ = renderers.render("deck.pptx", outline)
    assert pptx[:2] == b"PK"


def test_spreadsheet_formula_injection_is_neutralised():
    from openpyxl import load_workbook

    data, _ = renderers.render("t.xlsx", 'name,value\n=HYPERLINK("http://x"),-5\n')
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["A2"].value.startswith("'=")
    assert ws["B2"].value == -5
    csv_out, _ = renderers.render("t.csv", "@SUM(A1),1\n")
    assert csv_out.decode().startswith("'@SUM")


def test_registry_sync_and_usage():
    ctx = make_ctx()
    n = registry.sync(ctx.store)
    assert n == len(registry.load_skills())
    registry.record_use(ctx.store, "code", success=False)
    row = ctx.store.select("skills", {"id": "code"})[0]
    assert (row["usage_count"], row["success_count"]) == (1, 0)


# -- notifier ----------------------------------------------------------------


def test_deadline_alerts_only_for_active_events():
    ctx = make_ctx()
    soon = (FIXED_NOW + timedelta(hours=30)).isoformat()
    for i, status in enumerate(["building", "rejected"]):
        ctx.store.insert(
            "events",
            {
                "source_id": "d",
                "external_id": str(i),
                "title": status,
                "url": "https://x/",
                "status": status,
                "submission_deadline": soon,
            },
        )
    assert notifier.deadline_alerts(ctx) == 1
    notifier.deadline_alerts(ctx)
    assert len(ctx.store.select("alerts")) == 1  # deduplicated


def test_weekly_summary_includes_rejection_patterns():
    ctx = make_ctx()
    ctx.store.insert(
        "scores",
        {
            "event_id": "e",
            "passed_hard_filters": False,
            "hard_filters": [{"name": "solo_allowed", "passed": False}],
        },
    )
    ctx.store.insert("spend_log", {"cost_usd": Decimal("0.42"), "agent": "x", "model": "m"})
    ctx.clock = lambda: FIXED_NOW - timedelta(days=FIXED_NOW.weekday())  # a Monday
    stats = notifier.run(ctx)
    assert stats.get("weekly")
    body = ctx.store.select("alerts", {"kind": "weekly_summary"})[0]["body"]
    assert "hard filter: solo_allowed (1)" in body and "$0.42" in body
